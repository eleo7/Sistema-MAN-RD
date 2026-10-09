from io import BytesIO
from zipfile import ZipFile

import pytest
from PIL import Image
from pypdf import PdfReader
from docx import Document
from app import app, clean_image
from documents import RECIPIENT, ADDRESS, WEBSITE


def picture(color='green'):
    output = BytesIO()
    Image.new('RGB', (640, 320), color).save(output, 'PNG')
    output.seek(0)
    return output


def payload(count=1, output='pdf'):
    data = dict(number='042/2026', date='2026-10-09', city='Goiânia', clients='150', signer='Cesar Augusto Guarilha', role='Gerente de Manutenção e Automação RD Centro', model='palmeira', format=output, point_id=[f'p{i}' for i in range(count)], logo=(picture('blue'), 'logo.png'))
    for i in range(count):
        data[f'location_p{i}'] = f'Rua {i + 1}, Jardim Goiás, Goiânia'
        data[f'photo_p{i}'] = (picture(), 'evidencia.png')
    return data


@pytest.fixture
def client():
    app.config['TESTING'] = True
    return app.test_client()


def test_home(client):
    response = client.get('/')
    assert response.status_code == 200
    assert 'Manutenção RD - GYN' in response.text
    assert 'Leonardo de Souza Neres' in response.text


def test_one_page_pdf(client):
    response = client.post('/generate', data=payload())
    assert response.status_code == 200
    pdf = PdfReader(BytesIO(response.data))
    assert len(pdf.pages) == 1
    text = pdf.pages[0].extract_text()
    for expected in ['09 de outubro de 2026', 'Localização:', 'Evidência fotográfica', '150', 'Cesar Augusto Guarilha', 'EQUATORIAL ENERGIA GOIÁS', WEBSITE]:
        assert expected in text
    assert 'Ponto 1' not in text
    assert len(pdf.pages[0].images) == 2
    chunks = []
    pdf.pages[0].extract_text(visitor_text=lambda text, cm, tm, font, size: chunks.append((text, str(font.get('/BaseFont')) if font else '')))
    assert any('150' in text and 'Bold' in font for text, font in chunks)


def test_multipage_headers_and_footers(client):
    response = client.post('/generate', data=payload(7))
    assert response.status_code == 200
    pdf = PdfReader(BytesIO(response.data))
    assert len(pdf.pages) > 1
    for page in pdf.pages:
        assert 'EQUATORIAL ENERGIA GOIÁS' in page.extract_text()
        assert WEBSITE in page.extract_text()
        assert len(page.images) >= 1
    text = '\n'.join(p.extract_text() for p in pdf.pages)
    assert 'Ponto 7 — Localização:' in text
    assert 'Evidência fotográfica — Ponto 7' in text


def test_zip_word_content_and_bold(client):
    response = client.post('/generate', data=payload(2, 'zip'))
    assert response.status_code == 200
    with ZipFile(BytesIO(response.data)) as archive:
        assert set(archive.namelist()) == {'oficio-amma.pdf', 'oficio-amma.docx'}
        word = Document(BytesIO(archive.read('oficio-amma.docx')))
        pdf = PdfReader(BytesIO(archive.read('oficio-amma.pdf')))
    pdf_text = ' '.join(' '.join(p.extract_text().split()) for p in pdf.pages)
    for p in word.paragraphs:
        if p.text:
            assert ' '.join(p.text.split()) in pdf_text
    quantity = [run for p in word.paragraphs for run in p.runs if run.text == '150']
    assert quantity and quantity[0].bold is True
    assert word.styles['Normal'].font.name == 'Times New Roman'
    assert word.styles['Normal'].paragraph_format.line_spacing == 1.5
    section = word.sections[0]
    assert not section.different_first_page_header_footer
    assert 'graphic' in section.header._element.xml
    assert ADDRESS in section.footer.paragraphs[0].text
    assert WEBSITE in section.footer.paragraphs[0].text


def test_optional_fields_and_tree_model(client):
    data = payload(output='docx')
    data['clients'] = data['number'] = ''
    data['model'] = 'arvore'
    response = client.post('/generate', data=data)
    assert response.status_code == 200
    text = '\n'.join(p.text for p in Document(BytesIO(response.data)).paragraphs)
    assert 'Supressão de Árvore: com risco de queda ou contato com rede elétrica de distribuição' in text
    assert 'Ressalta-se' not in text
    assert 'Ofício nº' not in text


def test_point_validation(client):
    data = payload(2)
    del data['photo_p1']
    data['location_p0'] = ' '
    response = client.post('/generate', data=data)
    assert response.status_code == 422
    assert 'Ponto 2: anexe a foto da evidência.' in response.json['errors']
    assert any('Ponto 1: localização' in e for e in response.json['errors'])


@pytest.mark.parametrize('name,value', [('date', '2026-02-30'), ('clients', '-1'), ('clients', '1.5'), ('model', 'x'), ('format', 'exe'), ('city', ''), ('signer', ''), ('role', '')])
def test_invalid_fields(client, name, value):
    data = payload()
    data[name] = value
    assert client.post('/generate', data=data).status_code == 422


def test_invalid_photo(client):
    data = payload()
    data['photo_p0'] = (BytesIO(b'not an image'), 'photo.png')
    response = client.post('/generate', data=data)
    assert response.status_code == 422
    assert 'Ponto 1:' in response.json['errors'][0]


def test_logo_crop():
    image = Image.new('RGBA', (100, 100), (0, 0, 0, 0))
    image.paste((0, 100, 200, 255), (20, 30, 80, 70))
    source = BytesIO()
    image.save(source, 'PNG')
    source.seek(0)
    cropped = clean_image(source, logo=True)
    assert Image.open(BytesIO(cropped)).size == (60, 40)


def test_long_location_and_xml_escaping(client):
    data = payload(output='zip')
    data['location_p0'] = '<Praça & Rua> ' + 'Localização longa ' * 90
    data['signer'] = 'João & Maria <RD>'
    response = client.post('/generate', data=data)
    assert response.status_code == 200
    with ZipFile(BytesIO(response.data)) as archive:
        word = Document(BytesIO(archive.read('oficio-amma.docx')))
        assert any('João & Maria <RD>' == p.text for p in word.paragraphs)


def test_missing_logo(client, monkeypatch, tmp_path):
    import app as app_module
    monkeypatch.setattr(app_module, 'ROOT', tmp_path)
    data = payload()
    del data['logo']
    response = client.post('/generate', data=data)
    assert response.status_code == 422
    assert any(e.startswith('Logo:') for e in response.json['errors'])


def test_project_logo_is_automatic(client, monkeypatch, tmp_path):
    import app as app_module
    monkeypatch.setattr(app_module, 'ROOT', tmp_path)
    (tmp_path / 'assets').mkdir()
    (tmp_path / 'assets' / 'logo.png').write_bytes(picture('blue').getvalue())
    page = client.get('/').text
    assert 'Logo oficial configurado.' in page
    assert 'name="logo"' not in page
    data = payload(output='zip')
    del data['logo']
    response = client.post('/generate', data=data)
    assert response.status_code == 200
    with ZipFile(BytesIO(response.data)) as archive:
        pdf = PdfReader(BytesIO(archive.read('oficio-amma.pdf')))
        assert len(pdf.pages[0].images) == 2
        word = Document(BytesIO(archive.read('oficio-amma.docx')))
        assert 'graphic' in word.sections[0].header._element.xml


@pytest.mark.parametrize('count', [1, 2])
def test_updated_signature_and_locations(client, count):
    data = payload(count, 'zip')
    del data['signer']
    del data['role']
    response = client.post('/generate', data=data)
    assert response.status_code == 200
    with ZipFile(BytesIO(response.data)) as archive:
        word = Document(BytesIO(archive.read('oficio-amma.docx')))
        pdf = PdfReader(BytesIO(archive.read('oficio-amma.pdf')))
    word_text = '\n'.join(p.text for p in word.paragraphs)
    pdf_text = '\n'.join(p.extract_text() for p in pdf.pages)
    for text in [word_text, pdf_text]:
        normalized = ' '.join(text.split())
        assert 'processo de supressão de Palmeira Imperial' in normalized
        assert 'supressão de uma' not in normalized
        assert 'THIAGO DUTRA SILVA' in normalized
        assert 'ENGENHEIRO AGRÔNOMO E TECNÓL. EM GEOPROCESSAMENTO' in normalized
        assert 'CREA: Nº 17219/D-GO' in normalized
        assert 'COORDENADOR DE MEIO AMBIENTE E RESPONSÁVEL TÉCNICO' in normalized
        if count > 1:
            assert 'nas seguintes localizações:' in normalized
            assert 'Ponto 1 — Rua 1' in normalized
            assert 'Ponto 2 — Rua 2' in normalized
        else:
            assert 'na seguinte localização:' in normalized
    signature = next(p for p in word.paragraphs if p.text == 'THIAGO DUTRA SILVA')
    assert 'w:pBdr' in signature._p.xml
    assert signature.runs[0].bold
    assert signature.paragraph_format.keep_with_next
